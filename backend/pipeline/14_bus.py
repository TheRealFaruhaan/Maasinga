"""Step 14: Greater Male bus ridership from the organiser's MTCC file (SAMPLE data).

Input: data/raw/MTCC/MTCC.xlsx (sheets bus_monthly, bus_daily). Its own notes say: route names and codes are real,
but every passenger figure is randomly generated and is not MTCC data. The API returns these rows with sample: true.

Output: bus.json
- monthly: [month "YYYY-MM", route, paid, free of charge, total], Jan 2023 to Dec 2025, routes R1, R2, R3, R4, R7
- daily:   [date "YYYY-MM-DD", route, paid, free of charge, total], 2025, routes R1, R2, R4
- routes:  {code: name} as written in the file
"""
from openpyxl import load_workbook

from common import RAW, write_json

SRC = RAW / "MTCC" / "MTCC.xlsx"

if __name__ == "__main__":
    wb = load_workbook(SRC, read_only=True, data_only=True)
    monthly = [list(r) for r in wb["bus_monthly"].iter_rows(min_row=2, values_only=True) if r[0]]
    daily = [list(r) for r in wb["bus_daily"].iter_rows(min_row=2, values_only=True) if r[0]]
    write_json("bus.json", {
        "sample": True,
        "note": "Route names and codes are real; all passenger figures are randomly generated and are not actual MTCC data (MTCC file notes).",
        "routes": {r[3]: r[4] for r in monthly},
        "monthly": [[m, rid, paid, foc, total] for m, _, _, rid, _, paid, foc, total in monthly],
        "daily": [[str(d)[:10], rid, paid, foc, total] for d, _, rid, _, paid, foc, total in daily],
    })
    print(f"{len(monthly)} monthly rows, {len(daily)} daily rows")
